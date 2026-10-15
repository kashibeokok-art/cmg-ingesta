"""Los libros de Excel que entrega la app.

- **Fechas como fechas**, no texto: en Excel funcionan `=MES()`, `=AÑO()` y los
  filtros por fecha (pedido del usuario en el programa anterior).
- **Encabezados con su unidad** ("CMg [USD/MWh]") y el origen en palabras
  ("Definitivo", "Preliminar"): el archivo lo abre gente que no conoce la base.
- **Una hoja "Info"** con barra, rango, formulas y avisos (dias preliminares):
  el archivo tiene que explicarse solo meses despues.
- Si una serie no cabe en una hoja (1.048.575 filas), se parte por año.
"""

from collections.abc import Collection, Mapping
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_FILAS = 1_048_575

ORIGEN = {
    "pagina_cen_def": "Definitivo",
    "pagina_cen_pre": "Preliminar",
    "maestro_cmg_db": "Maestro (historico)",
}

#: Nombre interno -> encabezado en el Excel.
ENCABEZADOS = {
    "fecha": "Fecha",
    "hora": "Hora",
    "minuto": "Minuto",
    "fecha_hora": "Fecha y hora",
    "bloque": "Bloque",
    "cmg_usd_mwh": "CMg [USD/MWh]",
    "es_hora_extra": "Hora extra (abril)",
    "barra": "Barra",
    "origen": "Origen",
    "version": "Version",
    "usd_clp": "Tipo de cambio [CLP/USD]",
    "publicado": "Publicado",
    "cargado": "Cargado en la base",
    "periodo": "Periodo",
    "desde": "Desde",
    "hasta": "Hasta",
    "cuartos": "Cuartos de hora",
    "A": "A [USD/MWh]",
    "B": "B [USD/MWh]",
    "C": "C [USD/MWh]",
    "Solar": "Solar [USD/MWh]",
    "NoSolar": "NoSolar [USD/MWh]",
    "Total": "Total [USD/MWh]",
    "promedio": "Promedio [USD/MWh]",
    "desv_estandar": "Desv. estandar [USD/MWh]",
    "minimo": "Minimo [USD/MWh]",
    "maximo": "Maximo [USD/MWh]",
    "en_cero": "Cuartos en cero",
    "cmg_referencia": "CMg referencia [USD/MWh]",
    "cmg_comparada": "CMg comparada [USD/MWh]",
    "riesgo_usd_mwh": "Riesgo [USD/MWh]",
    "riesgo_pct": "Riesgo [%]",
    "riesgo_minimo": "Riesgo minimo [USD/MWh]",
    "riesgo_maximo": "Riesgo maximo [USD/MWh]",
    "referencia_en_cero": "Cuartos con referencia en cero",
    "antes": "Antes",
    "despues": "Despues",
    "documento": "Documento del Coordinador",
    "detectado": "Detectado",
    "cuartos_distintos": "Cuartos que cambiaron",
    "dif_max_usd": "Cambio maximo [USD/MWh]",
    "dif_prom_usd": "Cambio promedio [USD/MWh]",
}

_FECHA = "dd-mm-yyyy"
_FECHA_HORA = "dd-mm-yyyy hh:mm"
_USD = "#,##0.00"
_ENTERO = "#,##0"
_PCT = "0.00%"


def _formato(encabezado: str) -> str | None:
    if encabezado in ("Fecha", "Desde", "Hasta", "Detectado", "Publicado", "Cargado en la base"):
        return _FECHA
    if encabezado == "Fecha y hora":
        return _FECHA_HORA
    if encabezado == "Riesgo [%]":
        return _PCT
    if "[USD/MWh]" in encabezado or "[CLP/USD]" in encabezado:
        return _USD
    if encabezado.startswith("Cuartos"):
        return _ENTERO
    return None


def preparar(df: pd.DataFrame) -> pd.DataFrame:
    """Encabezados legibles y el origen en palabras. No cambia los valores."""
    salida = df.copy()
    if "origen" in salida:
        salida["origen"] = salida["origen"].map(lambda o: ORIGEN.get(str(o), str(o)))
    for col in ("antes", "despues"):
        if col in salida:
            salida[col] = salida[col].map(lambda o: ORIGEN.get(str(o), "" if o is None else str(o)))
    return salida.rename(columns=ENCABEZADOS)


def _dar_formato(hoja: Worksheet, usd: Collection[str] = ()) -> None:
    """Encabezado en negrita y fijo, anchos de columna y formato por columna.

    `usd`: columnas en USD/MWh cuyo nombre no lo dice (en la hoja ancha, cada
    columna se llama como su barra).
    """
    hoja.freeze_panes = "A2"
    relleno = PatternFill("solid", fgColor="DDE6F0")
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = relleno
    for i, celda in enumerate(hoja[1], start=1):
        encabezado = str(celda.value or "")
        letra = get_column_letter(i)
        hoja.column_dimensions[letra].width = max(10, min(40, len(encabezado) + 2))
        formato = _USD if encabezado in usd else _formato(encabezado)
        if formato:
            for (c,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                c.number_format = formato


def _hojas_de_serie(nombre: str, serie: pd.DataFrame) -> dict[str, pd.DataFrame]:
    if len(serie) <= MAX_FILAS:
        return {nombre: serie}
    anios = pd.to_datetime(serie["Fecha"]).dt.year
    return {f"{nombre} {a}": serie[anios == a] for a in sorted(anios.unique())}


def hoja_info(datos: Mapping[str, object]) -> pd.DataFrame:
    def texto(v: object) -> str:
        if isinstance(v, datetime):
            return f"{v:%d-%m-%Y %H:%M}"
        if isinstance(v, date):
            return f"{v:%d-%m-%Y}"
        return str(v)

    return pd.DataFrame({"Concepto": list(datos), "Valor": [texto(v) for v in datos.values()]})


def escribir(
    ruta: Path,
    hojas: Mapping[str, pd.DataFrame],
    usd: Mapping[str, Collection[str]] | None = None,
) -> Path:
    """Escribe el libro y le da formato. Las hojas vacias no se escriben."""
    with pd.ExcelWriter(ruta, engine="openpyxl") as libro:
        for nombre, df in hojas.items():
            if df.empty:
                continue
            df.to_excel(libro, sheet_name=nombre[:31], index=False)
            _dar_formato(libro.sheets[nombre[:31]], (usd or {}).get(nombre, ()))
    return ruta


# ------------------------------------------------------ varias barras, o todas

#: Encabezados de la parte fija de la serie ancha (una columna por barra).
ENCABEZADOS_ANCHA = {
    "fecha": "Fecha",
    "hora": "Hora",
    "minuto": "Minuto",
    "fecha_hora": "Fecha y hora",
    "bloque": "Bloque",
    "version": "Version",
    "usd_clp": "Tipo de cambio [CLP/USD]",
}


def libro_varias(
    ruta: Path,
    ancha: pd.DataFrame,
    barras: list[str],
    bloques: pd.DataFrame,
    versiones: pd.DataFrame,
    revisiones: pd.DataFrame,
    info: Mapping[str, object],
) -> Path:
    """Varias barras: el CMg cada 15 minutos con UNA COLUMNA POR BARRA, y los bloques."""
    hojas = {
        "CMg 15 min": ancha.rename(columns=ENCABEZADOS_ANCHA),
        "Bloques por mes": preparar(bloques),
        "Version por dia": preparar(versiones),
        "Revisiones": preparar(revisiones),
        "Info": hoja_info(info),
    }
    return escribir(ruta, hojas, usd={"CMg 15 min": set(barras)})


def libro_bloques_todas(
    ruta: Path,
    bloques: pd.DataFrame,
    versiones: pd.DataFrame,
    revisiones: pd.DataFrame,
    info: Mapping[str, object],
) -> Path:
    """Todas las barras: los bloques por mes (el dato cada 15 minutos va en el CSV)."""
    hojas = {
        "Bloques por mes": preparar(bloques),
        "Version por dia": preparar(versiones),
        "Revisiones": preparar(revisiones),
        "Info": hoja_info(info),
    }
    return escribir(ruta, hojas)


def agregar_csv_ancho(ruta: Path, ancha: pd.DataFrame, con_encabezado: bool) -> None:
    """Agrega un tramo de la serie ancha al CSV. `;` y coma decimal: Excel chileno lo
    abre con doble clic. Fechas como dd-mm-aaaa, que Excel reconoce como fecha."""
    tramo = ancha.copy()
    # 5 decimales, los que publica el Coordinador. El historico del Maestro venia en
    # 32 bits (107.11000061035156 por 107.11): redondear no pierde nada real, quita
    # ese ruido y achica el archivo a casi la mitad.
    barras = [c for c in tramo.columns if c not in ENCABEZADOS_ANCHA]
    tramo[barras] = tramo[barras].astype("float64").round(5)
    tramo["fecha"] = pd.to_datetime(tramo["fecha"]).dt.strftime("%d-%m-%Y")
    tramo["fecha_hora"] = pd.to_datetime(tramo["fecha_hora"]).dt.strftime("%d-%m-%Y %H:%M")
    tramo.rename(columns=ENCABEZADOS_ANCHA).to_csv(
        ruta,
        mode="w" if con_encabezado else "a",
        header=con_encabezado,
        index=False,
        sep=";",
        decimal=",",
        encoding="utf-8-sig" if con_encabezado else "utf-8",
    )


def libro_cmg(
    ruta: Path,
    serie: pd.DataFrame,
    bloques: pd.DataFrame,
    estadisticas: pd.DataFrame,
    versiones: pd.DataFrame,
    revisiones: pd.DataFrame,
    info: Mapping[str, object],
) -> Path:
    """CMg de una barra: la serie de 15 minutos, los bloques por mes y la Info."""
    hojas = {
        **_hojas_de_serie("CMg 15 min", preparar(serie)),
        "Bloques por mes": preparar(bloques),
        "Estadisticas por bloque": preparar(estadisticas),
        "Version por dia": preparar(versiones),
        "Revisiones": preparar(revisiones),
        "Info": hoja_info(info),
    }
    return escribir(ruta, hojas)


def libro_riesgo(
    ruta: Path,
    serie: pd.DataFrame,
    resumen: pd.DataFrame,
    versiones: pd.DataFrame,
    revisiones: pd.DataFrame,
    info: Mapping[str, object],
) -> Path:
    """Riesgo nodal: la serie de 15 minutos, el resumen por mes y bloque, y la Info."""
    hojas = {
        **_hojas_de_serie("Riesgo 15 min", preparar(serie)),
        "Resumen por mes y bloque": preparar(resumen),
        "Version por dia": preparar(versiones),
        "Revisiones": preparar(revisiones),
        "Info": hoja_info(info),
    }
    return escribir(ruta, hojas)
