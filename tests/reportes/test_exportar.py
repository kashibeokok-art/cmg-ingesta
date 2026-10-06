"""Tests de los exportes."""

from collections.abc import Callable
from datetime import datetime
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd
import pytest

from cmg_ingesta.reportes import exportar

# `con` y `sembrar` vienen de tests/conftest.py: pytest las inyecta por nombre.


Sembrar = Callable[..., None]


@pytest.fixture
def base(con: duckdb.DuckDBPyConnection, tmp_path: Path, sembrar: Sembrar) -> Path:
    silver = tmp_path / "silver"
    sembrar(con, silver, 2024, 6, "BARRA_1", dias=2)
    return silver


# ---------------------------------------------------------- nombre_unico


def test_nombre_unico_la_primera_vez(tmp_path: Path) -> None:
    assert exportar.nombre_unico(tmp_path, "informe", ".csv").name == "informe.csv"


def test_nombre_unico_versiona_sin_sobrescribir(tmp_path: Path) -> None:
    (tmp_path / "informe.csv").write_text("x", encoding="utf-8")
    assert exportar.nombre_unico(tmp_path, "informe", ".csv").name == "informe_v2.csv"

    (tmp_path / "informe_v2.csv").write_text("x", encoding="utf-8")
    assert exportar.nombre_unico(tmp_path, "informe", ".csv").name == "informe_v3.csv"


def test_nombre_unico_crea_la_carpeta(tmp_path: Path) -> None:
    destino = tmp_path / "nueva" / "subcarpeta"
    exportar.nombre_unico(destino, "x", ".csv")
    assert destino.is_dir()


def test_nombre_seguro_quita_caracteres_prohibidos() -> None:
    assert exportar._nombre_seguro('a<b>c:d"e/f\\g|h?i*j') == "a_b_c_d_e_f_g_h_i_j"


# ------------------------------------------------------------------- csv


def test_csv_usa_punto_y_coma_y_coma_decimal(tmp_path: Path) -> None:
    """Para que Excel en configuracion chilena lo abra en columnas."""
    df = pd.DataFrame({"a": [1.5], "b": ["x"]})
    ruta = exportar.escribir_csv(df, tmp_path / "t.csv")
    texto = ruta.read_text(encoding="utf-8-sig")
    assert "a;b" in texto
    assert "1,5;x" in texto


# ----------------------------------------------------------------- excel


def test_excel_crea_una_hoja_por_entrada(tmp_path: Path) -> None:
    hojas = {"Uno": pd.DataFrame({"x": [1]}), "Dos": pd.DataFrame({"y": [2]})}
    ruta, omitidas = exportar.escribir_excel(hojas, tmp_path / "t.xlsx")
    assert omitidas == []
    libro = openpyxl.load_workbook(ruta, read_only=True)
    assert libro.sheetnames == ["Uno", "Dos"]
    libro.close()


def test_excel_omite_e_informa_una_hoja_que_no_cabe(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Una hoja demasiado grande no revienta: se omite y se devuelve su nombre."""
    monkeypatch.setattr(exportar, "MAX_FILAS_EXCEL", 2)
    hojas = {"Chica": pd.DataFrame({"x": [1]}), "Grande": pd.DataFrame({"x": [1, 2, 3]})}
    _, omitidas = exportar.escribir_excel(hojas, tmp_path / "t.xlsx")
    assert omitidas == ["Grande"]


def test_hojas_por_año_separa_y_saca_las_columnas_de_particion() -> None:
    df = pd.DataFrame({"anio": [2023, 2024], "mes": [1, 1], "cmg_usd_mwh": [10.0, 20.0]})
    hojas = exportar.hojas_por_año(df)
    assert sorted(hojas) == ["2023", "2024"]
    assert "anio" not in hojas["2023"].columns


def test_hoja_info_tiene_dos_columnas() -> None:
    df = exportar.hoja_info({"barra": "X", "filas": 10})
    assert list(df.columns) == ["concepto", "valor"]
    assert len(df) == 2


# -------------------------------------------------------- exportar_cmg


def test_exportar_cmg_escribe_los_tres_formatos(
    con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path
) -> None:
    rutas = exportar.exportar_cmg(
        con,
        base,
        "BARRA_1",
        (2024, 6),
        (2024, 6),
        tmp_path / "out",
        formatos=("excel", "csv", "parquet"),
    )
    extensiones = sorted({r.suffix for r in rutas})
    assert extensiones == [".csv", ".parquet", ".xlsx"]
    assert all(r.exists() for r in rutas)


def test_el_excel_trae_la_hoja_del_año_bloques_e_info(
    con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path
) -> None:
    rutas = exportar.exportar_cmg(con, base, "BARRA_1", (2024, 6), (2024, 6), tmp_path / "out")
    libro = openpyxl.load_workbook(rutas[0], read_only=True)
    assert libro.sheetnames == ["2024", "Bloques", "Info"]
    libro.close()


def test_la_fecha_del_resumen_es_un_datetime_en_excel(
    con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path
) -> None:
    """Si fuera texto, en Excel no funcionarian =MES() ni =AÑO()."""
    rutas = exportar.exportar_cmg(con, base, "BARRA_1", (2024, 6), (2024, 6), tmp_path / "out")
    libro = openpyxl.load_workbook(rutas[0])
    hoja = libro["Bloques"]
    encabezados = [c.value for c in hoja[1]]
    col = encabezados.index("fecha") + 1
    assert isinstance(hoja.cell(row=2, column=col).value, datetime)
    libro.close()


def test_exportar_sin_datos_falla_con_mensaje_claro(
    con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path
) -> None:
    with pytest.raises(ValueError, match="no hay datos"):
        exportar.exportar_cmg(con, base, "NO_EXISTE", (2024, 6), (2024, 6), tmp_path / "out")


def test_formato_desconocido_falla(
    con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path
) -> None:
    with pytest.raises(ValueError, match="no soportado"):
        exportar.exportar_cmg(
            con,
            base,
            "BARRA_1",
            (2024, 6),
            (2024, 6),
            tmp_path / "out",
            formatos=("pdf",),  # type: ignore[arg-type]
        )


def test_sin_formatos_falla(con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="al menos un formato"):
        exportar.exportar_cmg(
            con, base, "BARRA_1", (2024, 6), (2024, 6), tmp_path / "out", formatos=()
        )


def test_exportar_dos_veces_no_sobrescribe(
    con: duckdb.DuckDBPyConnection, base: Path, tmp_path: Path
) -> None:
    salida = tmp_path / "out"
    uno = exportar.exportar_cmg(
        con, base, "BARRA_1", (2024, 6), (2024, 6), salida, formatos=("csv",)
    )
    dos = exportar.exportar_cmg(
        con, base, "BARRA_1", (2024, 6), (2024, 6), salida, formatos=("csv",)
    )
    assert uno[0] != dos[0]
    assert uno[0].exists() and dos[0].exists()
