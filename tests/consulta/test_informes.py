"""Los Excel que entrega la app, leidos de vuelta como los abriria el usuario."""

from datetime import date, datetime
from pathlib import Path

import openpyxl
import psycopg
import pytest

from cmg_ingesta.consulta import informes

Conn = psycopg.Connection[tuple[object, ...]]
REF, COMP = "A.BLANCAS_____013", "STA.ELVIRA____013"


def info(libro: openpyxl.Workbook) -> dict[str, str]:
    return {str(f[0]): str(f[1]) for f in libro["Info"].iter_rows(min_row=2, values_only=True)}


def test_excel_del_cmg_de_una_barra(lectura: Conn, tmp_path: Path) -> None:
    ruta = informes.cmg_barra(
        lectura, REF, date(2025, 3, 30), date(2025, 4, 2), tmp_path / "a.xlsx"
    )

    libro = openpyxl.load_workbook(ruta)
    assert libro.sheetnames == [
        "CMg 15 min",
        "Bloques por mes",
        "Estadisticas por bloque",
        "Version por dia",
        "Info",
    ]  # sin revisiones en el rango: la hoja no se escribe
    serie = libro["CMg 15 min"]
    assert [c.value for c in serie[1]] == [
        "Fecha", "Hora", "Minuto", "Fecha y hora", "Bloque", "CMg [USD/MWh]",
        "Hora extra (abril)", "Version", "Tipo de cambio [CLP/USD]",
    ]  # fmt: skip
    assert serie.max_row == 1 + 4 * 96
    fecha, fecha_hora, cmg = serie["A2"], serie["D2"], serie["F2"]
    assert isinstance(fecha.value, datetime) and fecha.number_format == "dd-mm-yyyy"
    assert fecha_hora.value == datetime(2025, 3, 30, 0, 0)
    assert cmg.number_format == "#,##0.00"
    assert serie.freeze_panes == "A2"
    assert serie["H2"].value == "Definitivo" and serie[f"H{serie.max_row}"].value == "Preliminar"

    bloques = libro["Bloques por mes"]
    assert [bloques.cell(r, 1).value for r in range(2, bloques.max_row + 1)] == [
        "2025-03",
        "2025-04",
        "Rango completo",
    ]
    i = info(libro)
    assert i["Barra"] == REF and i["Desde"] == "30-03-2025"
    assert "2 dia(s) con datos preliminares" in i["Datos preliminares"]


def test_excel_del_riesgo_nodal(lectura: Conn, tmp_path: Path) -> None:
    ruta = informes.riesgo_nodal(
        lectura, REF, COMP, date(2025, 3, 1), date(2025, 3, 3), tmp_path / "r.xlsx"
    )
    libro = openpyxl.load_workbook(ruta)
    assert libro.sheetnames[:2] == ["Riesgo 15 min", "Resumen por mes y bloque"]
    serie = libro["Riesgo 15 min"]
    encabezados = [c.value for c in serie[1]]
    assert "Riesgo [%]" in encabezados
    pct = serie.cell(2, encabezados.index("Riesgo [%]") + 1)
    assert pct.number_format == "0.00%"
    i = info(libro)
    assert i["Formula"] == "riesgo = CMg(comparada) - CMg(referencia)"
    assert i["Datos preliminares"].startswith("Ninguno")
    assert libro["Version por dia"].cell(3, 2).value == "Definitivo v2"  # el 2025-03-02


@pytest.mark.parametrize(
    ("desde", "hasta", "mensaje"),
    [
        (date(2025, 3, 5), date(2025, 3, 1), "posterior"),
        (date(2020, 1, 1), date(2020, 2, 1), "queda fuera"),
    ],
)
def test_rango_invalido(
    lectura: Conn, tmp_path: Path, desde: date, hasta: date, mensaje: str
) -> None:
    with pytest.raises(ValueError, match=mensaje):
        informes.cmg_barra(lectura, REF, desde, hasta, tmp_path / "x.xlsx")
    assert not (tmp_path / "x.xlsx").exists()


def test_barra_sin_datos(lectura: Conn, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="no tiene datos"):
        informes.cmg_barra(
            lectura, "NO_EXISTE", date(2025, 3, 1), date(2025, 3, 2), tmp_path / "x.xlsx"
        )


def test_riesgo_de_una_barra_consigo_misma(lectura: Conn, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="la misma"):
        informes.riesgo_nodal(
            lectura, REF, REF, date(2025, 3, 1), date(2025, 3, 2), tmp_path / "x.xlsx"
        )


def test_nombre_sugerido_valido_en_windows() -> None:
    assert informes.nombre_sugerido("CMg", "A/B:C", "20250301") == "CMg_A_B_C_20250301.xlsx"


# ------------------------------------------------------ varias barras, o todas


def test_excel_de_varias_barras_una_columna_por_barra(lectura: Conn, tmp_path: Path) -> None:
    barras = [COMP, REF]
    ruta = informes.cmg_varias(
        lectura, barras, date(2025, 3, 1), date(2025, 3, 2), tmp_path / "v.xlsx"
    )
    libro = openpyxl.load_workbook(ruta)
    hoja = libro["CMg 15 min"]
    assert [c.value for c in hoja[1]] == [
        "Fecha", "Hora", "Minuto", "Fecha y hora", "Bloque", "Version",
        "Tipo de cambio [CLP/USD]", REF, COMP,
    ]  # fmt: skip
    assert hoja.max_row == 1 + 2 * 96
    assert hoja["H2"].number_format == "#,##0.00"
    bloques = libro["Bloques por mes"]
    assert [c.value for c in bloques[1]][:2] == ["Barra", "Periodo"]
    assert info(libro)["Barras sin datos en el rango"] == "ninguna"


def test_varias_barras_necesita_al_menos_dos(lectura: Conn, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="al menos dos"):
        informes.cmg_varias(lectura, [REF], date(2025, 3, 1), date(2025, 3, 2), tmp_path / "x.xlsx")


def test_demasiados_datos_para_excel_sugiere_el_csv(
    lectura: Conn, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(informes, "MAX_CELDAS_EXCEL", 100)
    with pytest.raises(ValueError, match="Todas las barras"):
        informes.cmg_varias(
            lectura, [REF, COMP], date(2025, 3, 1), date(2025, 3, 2), tmp_path / "x.xlsx"
        )
    assert not (tmp_path / "x.xlsx").exists()


def test_todas_las_barras_csv_ancho_y_excel_de_bloques(lectura: Conn, tmp_path: Path) -> None:
    """Un rango que cruza dos meses: el CSV se escribe por partes, con UN encabezado."""
    avisos: list[str] = []
    csv, xlsx = informes.cmg_todas(
        lectura, date(2025, 3, 31), date(2025, 4, 1), tmp_path / "todas.csv", avisar=avisos.append
    )
    lineas = csv.read_text(encoding="utf-8-sig").splitlines()
    encabezado = lineas[0].split(";")
    assert encabezado[:7] == [
        "Fecha", "Hora", "Minuto", "Fecha y hora", "Bloque", "Version", "Tipo de cambio [CLP/USD]",
    ]  # fmt: skip
    assert encabezado[7:] == ["A.BLANCAS_____013", "QUELLON_______013", "STA.ELVIRA____013"]
    assert len(lineas) == 1 + 2 * 96
    assert sum(1 for x in lineas if x.startswith("Fecha;")) == 1
    primera = lineas[1].split(";")
    assert primera[0] == "31-03-2025" and primera[3] == "31-03-2025 00:00"
    assert "," in primera[7] or primera[7].isdigit()  # coma decimal
    decimales = [len(v.split(",")[1]) for x in lineas[1:] for v in x.split(";")[7:] if "," in v]
    assert max(decimales) <= 5  # sin el ruido de 32 bits del historico
    assert sum(1 for a in avisos if a.startswith("Bajando")) == 2
    libro = openpyxl.load_workbook(xlsx)
    assert libro.sheetnames[0] == "Bloques por mes"
    assert "todas (3)" in info(libro)["Barras"]
